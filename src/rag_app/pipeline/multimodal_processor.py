import base64
import io
import os
import logging
from PIL import Image
import requests
from typing import List, Dict, Any, Optional
import fitz  # PyMuPDF
import hashlib
import json
from pathlib import Path
from datetime import datetime

logger = logging.getLogger(__name__)

class MultimodalProcessor:
    """
    Handles image extraction and description for multimodal RAG.
    Integrates with OpenRouter free vision models for educational content analysis.
    
    Based on best practices from:
    - PyMuPDF documentation for robust image extraction
    - Multimodal RAG implementations from research papers
    - Production-ready error handling and caching
    """
    
    def __init__(self, openrouter_api_key: str):
        self.api_key = openrouter_api_key
        self.base_url = "https://openrouter.ai/api/v1"
        # Use Qwen2.5-VL-72B for best educational content understanding
        self.vision_model = os.getenv('VISION_MODEL', 'qwen/qwen2.5-vl-72b-instruct:free')
        self.max_images_per_doc = int(os.getenv('MAX_IMAGES_PER_DOCUMENT', '20'))
        
        # Cache directory for image descriptions
        self.cache_dir = Path('media/cache/image_descriptions')
        self.cache_dir.mkdir(parents=True, exist_ok=True)
        
        logger.info(f"MultimodalProcessor initialized with model: {self.vision_model}")
    
    def _get_cache_key(self, image_data: bytes, context: str = "") -> str:
        """Generate cache key for image description"""
        content = image_data + context.encode('utf-8')
        return hashlib.md5(content).hexdigest()
    
    def _get_cached_description(self, cache_key: str) -> Optional[str]:
        """Get cached image description if exists"""
        cache_file = self.cache_dir / f"{cache_key}.json"
        if cache_file.exists():
            try:
                with open(cache_file, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                    return data.get('description')
            except Exception as e:
                logger.warning(f"Error reading cache file {cache_file}: {e}")
        return None
    
    def _save_cached_description(self, cache_key: str, description: str):
        """Save image description to cache"""
        cache_file = self.cache_dir / f"{cache_key}.json"
        try:
            with open(cache_file, 'w', encoding='utf-8') as f:
                json.dump({
                    'description': description,
                    'timestamp': datetime.now().isoformat()
                }, f, ensure_ascii=False, indent=2)
        except Exception as e:
            logger.warning(f"Error saving to cache file {cache_file}: {e}")
    
    def extract_images_from_pdf(self, pdf_path: str) -> List[Dict]:
        """Extract images from PDF using PyMuPDF with improved error handling"""
        images = []
        doc = None
        
        try:
            doc = fitz.open(pdf_path)
            logger.info(f"Processing PDF with {len(doc)} pages: {pdf_path}")
            
            for page_num in range(len(doc)):
                if len(images) >= self.max_images_per_doc:
                    logger.info(f"Reached maximum images limit ({self.max_images_per_doc})")
                    break
                    
                page = doc.load_page(page_num)
                
                # Get images using the recommended PyMuPDF approach
                image_list = page.get_images(full=True)
                
                for img_index, img in enumerate(image_list):
                    if len(images) >= self.max_images_per_doc:
                        break
                        
                    try:
                        # Extract image using xref (cross-reference number)
                        xref = img[0]
                        
                        # Method 1: Extract image data directly (recommended for reliability)
                        img_dict = doc.extract_image(xref)
                        img_data = img_dict["image"]
                        img_ext = img_dict["ext"]
                        width = img_dict["width"]
                        height = img_dict["height"]
                        
                        # Skip very small images (likely decorative elements)
                        if width < 50 or height < 50:
                            logger.debug(f"Skipping small image: {width}x{height}")
                            continue
                        
                        # Convert to PIL Image to ensure format compatibility
                        try:
                            pil_image = Image.open(io.BytesIO(img_data))
                            
                            # Convert to RGB if necessary (remove alpha channel for vision models)
                            if pil_image.mode in ('RGBA', 'LA', 'P'):
                                pil_image = pil_image.convert('RGB')
                            
                            # Save as PNG for consistency
                            img_buffer = io.BytesIO()
                            pil_image.save(img_buffer, format='PNG')
                            processed_img_data = img_buffer.getvalue()
                            
                            # Get image bbox for location context
                            try:
                                bbox = page.get_image_bbox(img)
                            except:
                                bbox = None
                            
                            images.append({
                                'data': processed_img_data,
                                'page': page_num + 1,
                                'index': img_index,
                                'bbox': bbox,
                                'width': width,
                                'height': height,
                                'size_bytes': len(processed_img_data),
                                'original_format': img_ext,
                                'xref': xref
                            })
                            
                            logger.debug(f"Extracted image {len(images)} from page {page_num + 1} ({width}x{height}, {img_ext})")
                            
                        except Exception as e:
                            logger.warning(f"Error processing image data for xref {xref}: {e}")
                            continue
                        
                    except Exception as e:
                        logger.warning(f"Error extracting image {img_index} from page {page_num + 1}: {e}")
                        continue
            
            logger.info(f"Successfully extracted {len(images)} images from PDF")
            
        except Exception as e:
            logger.error(f"Error processing PDF {pdf_path}: {e}")
        finally:
            if doc:
                doc.close()
            
        return images
    
    def extract_images_from_docx(self, docx_path: str) -> List[Dict]:
        """Extract images from DOCX files with improved error handling"""
        images = []
        try:
            from docx import Document
            
            doc = Document(docx_path)
            logger.info(f"Processing DOCX: {docx_path}")
            
            # Extract images from document relationships
            for rel_id, rel in doc.part.rels.items():
                if "image" in rel.target_ref.lower():
                    if len(images) >= self.max_images_per_doc:
                        break
                        
                    try:
                        img_data = rel.target_part.blob
                        
                        # Process image similar to PDF
                        try:
                            pil_image = Image.open(io.BytesIO(img_data))
                            
                            # Skip very small images
                            if pil_image.width < 50 or pil_image.height < 50:
                                continue
                            
                            # Convert to RGB for consistency
                            if pil_image.mode in ('RGBA', 'LA', 'P'):
                                pil_image = pil_image.convert('RGB')
                            
                            # Save as PNG
                            img_buffer = io.BytesIO()
                            pil_image.save(img_buffer, format='PNG')
                            processed_img_data = img_buffer.getvalue()
                            
                            images.append({
                                'data': processed_img_data,
                                'page': 1,  # DOCX doesn't have clear page concept
                                'index': len(images),
                                'width': pil_image.width,
                                'height': pil_image.height,
                                'size_bytes': len(processed_img_data),
                                'filename': rel.target_ref,
                                'rel_id': rel_id
                            })
                            
                            logger.debug(f"Extracted image {len(images)} from DOCX ({pil_image.width}x{pil_image.height})")
                            
                        except Exception as e:
                            logger.warning(f"Error processing DOCX image data: {e}")
                            continue
                            
                    except Exception as e:
                        logger.warning(f"Error extracting image from DOCX: {e}")
                        continue
            
            logger.info(f"Extracted {len(images)} images from DOCX")
            
        except ImportError:
            logger.error("python-docx not installed. Cannot extract images from DOCX.")
        except Exception as e:
            logger.error(f"Error processing DOCX {docx_path}: {e}")
            
        return images
    
    def extract_images_from_pptx(self, pptx_path: str) -> List[Dict]:
        """Extract images from PPTX files with improved error handling"""
        images = []
        try:
            from pptx import Presentation
            
            prs = Presentation(pptx_path)
            logger.info(f"Processing PPTX with {len(prs.slides)} slides: {pptx_path}")
            
            for slide_num, slide in enumerate(prs.slides):
                if len(images) >= self.max_images_per_doc:
                    break
                
                # Get slide title if available
                slide_title = ""
                try:
                    if slide.shapes.title and slide.shapes.title.text:
                        slide_title = slide.shapes.title.text.strip()
                    else:
                        slide_title = f"Slide {slide_num + 1}"
                except:
                    slide_title = f"Slide {slide_num + 1}"
                    
                for shape in slide.shapes:
                    if len(images) >= self.max_images_per_doc:
                        break
                        
                    # Check if shape contains an image
                    if hasattr(shape, 'image') and shape.image:
                        try:
                            img_data = shape.image.blob
                            
                            # Process image
                            try:
                                pil_image = Image.open(io.BytesIO(img_data))
                                
                                # Skip very small images
                                if pil_image.width < 50 or pil_image.height < 50:
                                    continue
                                
                                # Convert to RGB for consistency
                                if pil_image.mode in ('RGBA', 'LA', 'P'):
                                    pil_image = pil_image.convert('RGB')
                                
                                # Save as PNG
                                img_buffer = io.BytesIO()
                                pil_image.save(img_buffer, format='PNG')
                                processed_img_data = img_buffer.getvalue()
                                
                                images.append({
                                    'data': processed_img_data,
                                    'page': slide_num + 1,
                                    'index': len(images),
                                    'width': pil_image.width,
                                    'height': pil_image.height,
                                    'size_bytes': len(processed_img_data),
                                    'slide_title': slide_title,
                                    'shape_id': shape.shape_id if hasattr(shape, 'shape_id') else None
                                })
                                
                                logger.debug(f"Extracted image {len(images)} from slide {slide_num + 1} ({pil_image.width}x{pil_image.height})")
                                
                            except Exception as e:
                                logger.warning(f"Error processing PPTX image data: {e}")
                                continue
                                
                        except Exception as e:
                            logger.warning(f"Error extracting image from slide {slide_num + 1}: {e}")
                            continue
            
            logger.info(f"Extracted {len(images)} images from PPTX")
            
        except ImportError:
            logger.error("python-pptx not installed. Cannot extract images from PPTX.")
        except Exception as e:
            logger.error(f"Error processing PPTX {pptx_path}: {e}")
            
        return images
    
    def describe_image(self, image_data: bytes, context: str = "", page_info: str = "") -> str:
        """Convert image to educational description using vision model"""
        
        # Check cache first
        cache_key = self._get_cache_key(image_data, context)
        cached_desc = self._get_cached_description(cache_key)
        if cached_desc:
            logger.debug("Using cached image description")
            return cached_desc
        
        try:
            # Validate image data
            if not image_data or len(image_data) == 0:
                return "[IMAGE DESCRIPTION UNAVAILABLE - No image data]"
            
            # Convert image to base64
            image_b64 = base64.b64encode(image_data).decode('utf-8')
            
            # Create educational prompt optimized for learning
            prompt = f"""You are an educational AI assistant analyzing a document image. Provide a detailed description focusing on:

1. **Text Content**: Any text, formulas, equations, or mathematical expressions visible
2. **Visual Elements**: Charts, graphs, diagrams, tables, and their data/meaning
3. **Educational Concepts**: Theories, principles, or concepts being illustrated
4. **Relationships**: Connections between visual elements and their significance
5. **Processes**: Any step-by-step workflows or procedures shown
6. **Academic Content**: Scientific, technical, or scholarly information

Context: {context[:300] if context else 'Educational document analysis'}
Location: {page_info if page_info else 'Document image'}

Provide a comprehensive description that helps students understand and learn from this visual content. Be specific about data, formulas, and key concepts."""
            
            # Make API request to OpenRouter
            response = requests.post(
                f"{self.base_url}/chat/completions",
                headers={
                    "Authorization": f"Bearer {self.api_key}",
                    "Content-Type": "application/json",
                    "HTTP-Referer": "https://edumentorai.com",
                    "X-Title": "EduMentorAI Multimodal Processing"
                },
                json={
                    "model": self.vision_model,
                    "messages": [
                        {
                            "role": "user",
                            "content": [
                                {"type": "text", "text": prompt},
                                {
                                    "type": "image_url",
                                    "image_url": {
                                        "url": f"data:image/png;base64,{image_b64}"
                                    }
                                }
                            ]
                        }
                    ],
                    "max_tokens": 1500,
                    "temperature": 0.1  # Very low temperature for consistent educational descriptions
                },
                timeout=60  # Increased timeout for vision models
            )
            
            if response.status_code == 200:
                result = response.json()
                if 'choices' in result and len(result['choices']) > 0:
                    description = result['choices'][0]['message']['content']
                    
                    # Cache the description
                    self._save_cached_description(cache_key, description)
                    
                    logger.info("Successfully generated image description")
                    return description
                else:
                    logger.error(f"Invalid response format from OpenRouter: {result}")
                    return "[IMAGE DESCRIPTION UNAVAILABLE - Invalid API Response]"
            else:
                logger.error(f"OpenRouter API error: {response.status_code} - {response.text}")
                return f"[IMAGE DESCRIPTION UNAVAILABLE - API Error {response.status_code}]"
                
        except requests.exceptions.Timeout:
            logger.error("Timeout while calling vision model API")
            return "[IMAGE DESCRIPTION UNAVAILABLE - API Timeout]"
        except requests.exceptions.ConnectionError:
            logger.error("Connection error while calling vision model API")
            return "[IMAGE DESCRIPTION UNAVAILABLE - Connection Error]"
        except Exception as e:
            logger.error(f"Error describing image: {e}")
            return f"[IMAGE DESCRIPTION UNAVAILABLE - Error: {str(e)}]"
    
    def process_document_images(self, file_path: str, document_context: str = "") -> List[Dict]:
        """Process all images in a document and return descriptions"""
        file_ext = Path(file_path).suffix.lower()
        images = []
        
        logger.info(f"Processing images from {file_path} (type: {file_ext})")
        
        # Validate file exists
        if not os.path.exists(file_path):
            logger.error(f"File not found: {file_path}")
            return []
        
        # Extract images based on file type
        if file_ext == '.pdf':
            images = self.extract_images_from_pdf(file_path)
        elif file_ext == '.docx':
            images = self.extract_images_from_docx(file_path)
        elif file_ext == '.pptx':
            images = self.extract_images_from_pptx(file_path)
        else:
            logger.warning(f"Unsupported file type for image extraction: {file_ext}")
            return []
        
        if not images:
            logger.info("No images found in document")
            return []
        
        # Generate descriptions for extracted images
        described_images = []
        for i, img in enumerate(images):
            try:
                page_info = f"Page {img['page']}" if 'page' in img else f"Image {i+1}"
                if 'slide_title' in img and img['slide_title']:
                    page_info += f" ({img['slide_title']})"
                
                logger.info(f"Processing image {i+1}/{len(images)} from {page_info}")
                
                description = self.describe_image(
                    img['data'],
                    context=document_context,
                    page_info=page_info
                )
                
                described_images.append({
                    'page': img.get('page', 1),
                    'index': img.get('index', i),
                    'description': description,
                    'size_bytes': img.get('size_bytes', 0),
                    'width': img.get('width'),
                    'height': img.get('height'),
                    'slide_title': img.get('slide_title', ''),
                    'filename': img.get('filename', ''),
                    'bbox': img.get('bbox'),
                    'original_format': img.get('original_format', 'unknown'),
                    'processing_success': not description.startswith('[IMAGE DESCRIPTION UNAVAILABLE')
                })
                
                logger.info(f"Successfully processed image {i+1}/{len(images)} from {page_info}")
                
            except Exception as e:
                logger.error(f"Error processing image {i+1}: {e}")
                # Still add the image with error info
                described_images.append({
                    'page': img.get('page', 1),
                    'index': img.get('index', i),
                    'description': f"[IMAGE PROCESSING ERROR: {str(e)}]",
                    'size_bytes': img.get('size_bytes', 0),
                    'width': img.get('width'),
                    'height': img.get('height'),
                    'processing_success': False
                })
                continue
        
        success_count = sum(1 for img in described_images if img.get('processing_success', False))
        logger.info(f"Successfully processed {success_count}/{len(described_images)} images")
        return described_images
