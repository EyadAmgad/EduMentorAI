#!/usr/bin/env python3
"""
Integration test: Generate a real PowerPoint with images
"""

import os
import sys
import django

# Configure Django
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'rag_django.settings')
django.setup()

from django.contrib.auth import get_user_model
from rag_app.pipeline.SlideProcessor import SlideProcessor
from io import BytesIO

def test_full_presentation():
    """Test generating a complete presentation with images"""
    print("=" * 70)
    print("INTEGRATION TEST: Generate PowerPoint with Internet Images")
    print("=" * 70)
    
    processor = SlideProcessor()
    
    # Create a fake text file to process
    fake_content = """
    Machine Learning Overview
    
    Machine learning is a subset of artificial intelligence that enables 
    computers to learn from data without being explicitly programmed.
    
    Neural Networks
    
    Neural networks are computing systems inspired by biological neural 
    networks. They consist of interconnected layers of nodes.
    
    Deep Learning Applications
    
    Deep learning has revolutionized computer vision, natural language 
    processing, and autonomous systems.
    
    Computer Vision
    
    Computer vision enables machines to interpret and understand visual 
    information from the world.
    
    Natural Language Processing
    
    NLP allows computers to understand, interpret, and generate human language.
    """
    
    # Create a fake file object
    class FakeFile:
        def __init__(self, content, name):
            self.content = content
            self.name = name
            self.file = BytesIO(content.encode('utf-8'))
        
        def read(self):
            return self.content.encode('utf-8')
    
    fake_file = FakeFile(fake_content, "test_document.txt")
    
    # Get or create a test user
    User = get_user_model()
    user, _ = User.objects.get_or_create(
        username='test_user',
        defaults={'email': 'test@example.com'}
    )
    
    print("\n📝 Test Parameters:")
    print(f"   Files: 1 text file")
    print(f"   Slide count: 5")
    print(f"   Template: professional")
    print(f"   User: {user.username}")
    
    print("\n🚀 Generating presentation...")
    
    try:
        result = processor.generate_slides(
            files=[fake_file],
            slide_count=5,
            template='professional',
            title='Machine Learning Introduction',
            language='English',
            instructions='Create an educational presentation about machine learning',
            user=user,
            background_image=None,
            documents=None
        )
        
        if result['success']:
            print("\n✅ SUCCESS! Presentation generated")
            print(f"   File path: {result['filepath']}")
            
            # Check if file exists
            if os.path.exists(result['filepath']):
                file_size = os.path.getsize(result['filepath']) / 1024
                print(f"   File size: {file_size:.1f} KB")
                
                # Load and inspect the presentation
                from pptx import Presentation
                prs = Presentation(result['filepath'])
                
                print(f"\n📊 Presentation Analysis:")
                print(f"   Total slides: {len(prs.slides)}")
                
                slides_with_images = 0
                for i, slide in enumerate(prs.slides):
                    # Count images (pictures) in the slide
                    images_in_slide = sum(1 for shape in slide.shapes if shape.shape_type == 13)  # 13 = PICTURE
                    if images_in_slide > 0:
                        slides_with_images += 1
                        print(f"   Slide {i+1}: {images_in_slide} image(s) ✓")
                
                print(f"\n   Slides with images: {slides_with_images}")
                
                if slides_with_images > 0:
                    print("\n🎉 IMAGES SUCCESSFULLY ADDED TO PRESENTATION!")
                else:
                    print("\n⚠️  No images found in presentation")
                
            else:
                print(f"   ❌ File not found at path!")
                
        else:
            print(f"\n❌ FAILED: {result.get('error', 'Unknown error')}")
            
    except Exception as e:
        print(f"\n❌ Exception: {str(e)}")
        import traceback
        traceback.print_exc()
    
    print("\n" + "=" * 70)

if __name__ == '__main__':
    test_full_presentation()
